import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from china_trial_demo.database import connect
from china_trial_demo.importer import import_workbook


ROOT = Path(__file__).resolve().parents[1]


class ImporterTests(unittest.TestCase):
    def test_single_page_v4_round_trip(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = ROOT / "outputs" / "019ff196-7dcc-7242-8a16-64720facd020" / "合作方单试验填写表_v4.xlsx"
            workbook = load_workbook(source)
            sheet = workbook["单试验填写表"]
            values = {
                "F8": "V4-ORG", "B9": "第4版测试机构", "F9": "合作中", "B10": "V4-批次-001",
                "F10": "2026-09-01", "B11": "V4-TRIAL-001", "F11": "有效", "B15": "ChiCTR",
                "F15": "ChiCTR2600000002", "B16": "ChiCTR2600000002", "B18": "单页肺癌试验",
                "B24": "招募中", "F24": "干预性研究", "B26": "非小细胞肺癌", "B39": "中国",
                "F39": "2026-09-01", "B40": "第4版测试机构", "B41": "https://www.chictr.org.cn/showproj.html?proj=2",
                "A65": "中心-01", "B65": "测试医院", "C65": "中国｜上海市｜上海市", "D65": "招募中", "G65": "合作方报告中心",
                "A75": "入选", "B75": "ALL", "C75": "INC-001", "D75": "年龄", "E75": "年龄18至75岁", "F75": "是",
                "A76": "排除", "B76": "ALL", "C76": "EXC-001", "D76": "合并症", "E76": "未控制的活动性感染", "F76": "是",
            }
            for cell, value in values.items():
                sheet[cell] = value
            filled = root / "single-filled.xlsx"
            workbook.save(filled); workbook.close()
            db = root / "single.db"
            counts = import_workbook(db, filled)
            self.assertEqual(counts["试验"], 1)
            self.assertEqual(counts["研究中心"], 1)
            self.assertEqual(counts["入选标准"], 1)
            self.assertEqual(counts["排除标准"], 1)
            with connect(db, readonly=True) as connection:
                trial = dict(connection.execute("SELECT partner_trial_id,chictr_registration_number,countries FROM trials").fetchone())
                self.assertEqual(trial["partner_trial_id"], "V4-TRIAL-001")
                self.assertEqual(trial["chictr_registration_number"], "ChiCTR2600000002")
                self.assertEqual(trial["countries"], "中国")

    def test_generated_template_round_trip(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = ROOT / "outputs" / "019ff196-7dcc-7242-8a16-64720facd020" / "合作方临床试验数据填写模板_v3.xlsx"
            workbook = load_workbook(source)

            def write(sheet_name, row):
                sheet = workbook[sheet_name]
                headers = [cell.value for cell in sheet[3]]
                for key, value in row.items():
                    sheet.cell(row=5, column=headers.index(key) + 1, value=value)

            organization = "TEST-ORG"
            trial = "TRIAL-001"
            write("试验主表", {
                "模板版本": "中国临床试验合作方模板-第3版", "合作机构编码": organization, "合作机构名称": "测试机构",
                "合作状态": "合作中", "提交批次号": "测试批次-001", "数据截至日期": "2026-08-27",
                "合作方试验编号": trial, "ChiCTR注册号": "ChiCTR2600000001", "试验公开名称": "测试肺癌试验",
                "招募状态": "招募中", "研究类型": "干预性研究", "主要疾病或研究人群": "非小细胞肺癌",
                "主办单位": "测试机构", "组长单位": "测试医院", "国家或地区": "中国",
                "来源链接": "https://www.chictr.org.cn/showproj.html?proj=1", "最后核验日期": "2026-08-27", "记录状态": "有效",
            })
            write("研究中心", {"合作方试验编号": trial, "中心编号": "中心-01", "中心名称": "测试医院", "国家": "中国", "省": "上海市", "市": "上海市", "招募状态": "招募中", "证据类型": "合作方报告中心"})
            write("入排标准", {"合作方试验编号": trial, "标准类型": "入选", "队列编号": "ALL", "标准编号": "INC-001", "序号": 1, "类别": "age", "标准原文": "年龄18至75岁", "是否必需": "是", "患者字段": "age", "比较关系": "between", "比较值": "18|75", "单位": "year"})
            sheet = workbook["入排标准"]; headers = [cell.value for cell in sheet[3]]
            for key, value in {"合作方试验编号": trial, "标准类型": "排除", "队列编号": "ALL", "标准编号": "EXC-001", "序号": 1, "类别": "comorbidity", "标准原文": "未控制的活动性感染", "是否必需": "是"}.items():
                sheet.cell(row=6, column=headers.index(key) + 1, value=value)
            filled = root / "filled.xlsx"
            workbook.save(filled)
            workbook.close()
            db = root / "imported.db"
            counts = import_workbook(db, filled)
            self.assertEqual(counts["试验"], 1)
            self.assertEqual(counts["研究中心"], 1)
            with connect(db, readonly=True) as connection:
                self.assertEqual(connection.execute("SELECT COUNT(*) FROM organizations").fetchone()[0], 1)
                self.assertEqual(connection.execute("SELECT COUNT(*) FROM trials").fetchone()[0], 1)
                self.assertEqual(connection.execute("SELECT COUNT(*) FROM eligibility_criteria").fetchone()[0], 2)


if __name__ == "__main__":
    unittest.main()
