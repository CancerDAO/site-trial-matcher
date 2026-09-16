from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .database import connect, initialize, replace_trial_search
from .normalization import clean_text
from .workbook_v3 import MAIN_SHEET as V3_MAIN_SHEET, import_workbook_v3
from .workbook_v4 import SHEET_NAME as V4_SHEET_NAME, import_workbook_v4

TEMPLATE_VERSION = "中国试验合作方模板-第2版"
SHEET_NAME = "试验信息"
REQUIRED = ["模板版本", "合作机构编码", "合作机构名称", "合作状态", "提交批次号", "数据截至日期", "合作方试验编号",
            "试验中文名称", "招募状态", "研究类型", "主要疾病", "主办单位", "组长单位", "实施单位与地点（每条一行）",
            "入选标准（每条一行）", "排除标准（每条一行）", "来源链接", "最后核验日期", "记录状态"]
TRIAL_COLUMNS = {
    "合作方试验编号": "partner_trial_id", "方案编号": "protocol_number", "ChiCTR注册号": "chictr_registration_number",
    "其他注册号": "other_registry_ids", "试验中文名称": "title_zh", "试验英文名称": "title_en", "招募状态": "recruitment_status",
    "研究类型": "study_type", "研究阶段": "study_phase", "主要疾病": "primary_disease", "疾病别名": "disease_aliases",
    "病理类型": "histology", "疾病分期或状态": "disease_stage", "关键分子标志物": "biomarker_summary",
    "干预概述": "intervention_summary", "计划入组人数": "target_enrollment", "性别限制": "sex", "最小年龄（岁）": "minimum_age_years",
    "最大年龄（岁）": "maximum_age_years", "最大ECOG评分": "maximum_ecog", "开始日期": "first_enrollment_date",
    "计划结束日期": "planned_end_date", "主办单位": "primary_sponsor", "组长单位": "lead_institution", "来源链接": "source_url",
    "最后核验日期": "last_verified_date", "记录状态": "record_status", "研究摘要": "brief_summary", "数据备注": "data_notes",
}
CATEGORY_MAP = {"疾病":"disease", "病理":"histology", "分期":"stage", "生物标志物":"biomarker", "年龄":"age", "性别":"sex",
                "ECOG":"ecog", "既往治疗":"prior_therapy", "治疗线数":"treatment_line", "器官功能":"organ_function",
                "实验室":"lab", "合并症":"comorbidity", "妊娠":"pregnancy", "其他":"other"}
FIELD_MAP = {"年龄":"age", "性别":"sex", "ECOG":"ecog", "疾病":"cancer_type", "病理":"histology", "分期":"disease_stage",
             "生物标志物":"biomarkers", "突变":"biomarkers", "既往治疗":"prior_therapies", "治疗线数":"treatment_lines_completed",
             "合并症":"comorbidities", "妊娠":"pregnant"}
OPERATOR_MAP = {"等于":"eq", "不等于":"neq", "属于":"in", "不属于":"not_in", "大于等于":"gte", "小于等于":"lte",
                "介于":"between", "包含":"contains", "不包含":"not_contains", "已提供":"exists", "未提供":"not_exists"}


def _lines(value: Any) -> list[str]:
    return [line.strip() for line in str(value or "").replace("\r", "\n").split("\n") if line.strip()]


def _parts(line: str, expected: int, label: str) -> list[str]:
    parts = [part.strip() for part in line.split("｜")]
    if len(parts) != expected:
        raise ValueError(f"{label}格式错误：应有{expected}段并使用全角分隔符｜：{line}")
    return parts


def _read_rows(workbook: Any) -> list[dict[str, Any]]:
    if SHEET_NAME not in workbook.sheetnames:
        raise ValueError(f"缺少工作表：{SHEET_NAME}")
    sheet = workbook[SHEET_NAME]; headers = [clean_text(cell.value) for cell in sheet[3]]
    missing = [name for name in REQUIRED if name not in headers]
    if missing: raise ValueError(f"试验信息缺少列：{missing}")
    rows: list[dict[str, Any]] = []
    for row_number, values in enumerate(sheet.iter_rows(min_row=4, values_only=True), start=4):
        row = {headers[index]: values[index] for index in range(min(len(headers), len(values)))}
        if not any(value not in (None, "") for value in row.values()): continue
        if clean_text(row.get("合作机构编码")) == "示例": continue
        row["_row_number"] = row_number; rows.append(row)
    return rows


def _criterion_specs(value: Any, label: str) -> dict[int, dict[str, str]]:
    result: dict[int, dict[str, str]] = {}
    for line in _lines(value):
        order_text, category, field, operator, target, unit = _parts(line, 6, label)
        try: order = int(order_text)
        except ValueError as error: raise ValueError(f"{label}原文序号必须是整数：{line}") from error
        if category not in CATEGORY_MAP or field not in FIELD_MAP or operator not in OPERATOR_MAP:
            raise ValueError(f"{label}包含不支持的类别、患者字段或比较关系：{line}")
        if order in result: raise ValueError(f"{label}原文序号重复：{order}")
        result[order] = {"criterion_category": CATEGORY_MAP[category], "structured_field": FIELD_MAP[field],
                         "operator": OPERATOR_MAP[operator], "structured_value": target, "unit": unit}
    return result


def import_workbook(db_path: str | Path, workbook_path: str | Path) -> dict[str, int]:
    initialize(db_path); workbook = load_workbook(workbook_path, read_only=True, data_only=False)
    if V4_SHEET_NAME in workbook.sheetnames:
        try:
            return import_workbook_v4(db_path, workbook_path, workbook)
        finally:
            workbook.close()
    if V3_MAIN_SHEET in workbook.sheetnames:
        try:
            return import_workbook_v3(db_path, workbook_path, workbook)
        finally:
            workbook.close()
    rows = _read_rows(workbook); workbook.close()
    if not rows: raise ValueError("试验信息没有可导入的数据行")
    for row in rows:
        missing = [name for name in REQUIRED if row.get(name) in (None, "")]
        if missing: raise ValueError(f"试验信息第{row['_row_number']}行缺少必填值：{missing}")
        if clean_text(row["模板版本"]) != TEMPLATE_VERSION: raise ValueError(f"模板版本必须为{TEMPLATE_VERSION}")
        if "chictr.org.cn" not in clean_text(row["来源链接"]).casefold():
            raise ValueError(f"第{row['_row_number']}行来源链接必须是ChiCTR链接")
    identities = {(clean_text(row["合作机构编码"]), clean_text(row["合作机构名称"]), clean_text(row["合作状态"]),
                   clean_text(row["提交批次号"]), str(row["数据截至日期"])) for row in rows}
    if len(identities) != 1: raise ValueError("同一工作簿中的合作机构、合作状态、提交批次和数据截至日期必须一致")
    keys = [clean_text(row["合作方试验编号"]) for row in rows]
    if len(keys) != len(set(keys)): raise ValueError("合作方试验编号重复")
    organization_code, organization_name, partner_status, batch_name, as_of = next(iter(identities))

    counts = {"试验": 0, "实施单位": 0, "队列": 0, "入选标准": 0, "排除标准": 0}
    with connect(db_path) as connection:
        connection.execute("""INSERT INTO organizations(organization_code,organization_name_zh,partner_status,country_region)
            VALUES(?,?,?,'中国') ON CONFLICT(organization_code) DO UPDATE SET organization_name_zh=excluded.organization_name_zh,
            partner_status=excluded.partner_status,updated_at=CURRENT_TIMESTAMP""", (organization_code, organization_name, partner_status))
        organization_id = connection.execute("SELECT id FROM organizations WHERE organization_code=?", (organization_code,)).fetchone()[0]
        connection.execute("""INSERT INTO import_batches(organization_id,submission_batch_id,data_as_of_date,template_version,source_file,row_counts_json)
            VALUES(?,?,?,?,?,'{}') ON CONFLICT(organization_id,submission_batch_id) DO UPDATE SET data_as_of_date=excluded.data_as_of_date,
            template_version=excluded.template_version,source_file=excluded.source_file,imported_at=CURRENT_TIMESTAMP""",
            (organization_id, batch_name, as_of, TEMPLATE_VERSION, str(Path(workbook_path).resolve())))
        batch_id = connection.execute("SELECT id FROM import_batches WHERE organization_id=? AND submission_batch_id=?", (organization_id, batch_name)).fetchone()[0]
        for row in rows:
            columns = list(TRIAL_COLUMNS.values()); values = [row.get(header) for header in TRIAL_COLUMNS]
            assignments = ",".join(f"{name}=excluded.{name}" for name in columns[1:])
            connection.execute(f"INSERT INTO trials(organization_id,import_batch_id,{','.join(columns)}) VALUES(?,?,{','.join('?' for _ in columns)}) "
                f"ON CONFLICT(organization_id,partner_trial_id) DO UPDATE SET import_batch_id=excluded.import_batch_id,{assignments},updated_at=CURRENT_TIMESTAMP",
                [organization_id, batch_id, *values])
            trial_id = connection.execute("SELECT id FROM trials WHERE organization_id=? AND partner_trial_id=?", (organization_id, clean_text(row["合作方试验编号"]))).fetchone()[0]
            connection.execute("DELETE FROM trial_sites WHERE trial_id=?", (trial_id,)); connection.execute("DELETE FROM trial_cohorts WHERE trial_id=?", (trial_id,)); connection.execute("DELETE FROM eligibility_criteria WHERE trial_id=?", (trial_id,))
            for site_index, line in enumerate(_lines(row["实施单位与地点（每条一行）"]), start=1):
                site_id, name, province, city, status = _parts(line, 5, "实施单位与地点")
                connection.execute("INSERT INTO trial_sites(trial_id,site_id,site_name,province,city,recruitment_status,last_verified_date,record_status) VALUES(?,?,?,?,?,?,?,'有效')",
                    (trial_id, site_id or f"中心-{site_index:03d}", name, province, city, status, row["最后核验日期"])); counts["实施单位"] += 1
            for cohort_index, line in enumerate(_lines(row.get("队列与干预（每条一行）")), start=1):
                cohort_id, name, disease, biomarker, intervention, status = _parts(line, 6, "队列与干预")
                connection.execute("INSERT INTO trial_cohorts(trial_id,cohort_id,cohort_name,cohort_status,disease_scope,biomarker_scope,intervention_name,record_status) VALUES(?,?,?,?,?,?,?,'有效')",
                    (trial_id, cohort_id or f"队列-{cohort_index:03d}", name, status, disease, biomarker, intervention)); counts["队列"] += 1
            for kind, raw_header, structured_header, prefix, count_key in (("inclusion", "入选标准（每条一行）", "可计算入选条件（每条一行）", "入选", "入选标准"), ("exclusion", "排除标准（每条一行）", "可计算排除条件（每条一行）", "排除", "排除标准")):
                raw_lines = _lines(row[raw_header]); specs = _criterion_specs(row.get(structured_header), structured_header)
                if any(order < 1 or order > len(raw_lines) for order in specs): raise ValueError(f"{structured_header}的原文序号超出范围")
                for order, text in enumerate(raw_lines, start=1):
                    spec = specs.get(order, {}); category = spec.get("criterion_category", "other")
                    connection.execute("""INSERT INTO eligibility_criteria(trial_id,cohort_id,criterion_type,criterion_id,criterion_order,criterion_category,
                        criterion_text_zh,is_mandatory,structured_field,operator,structured_value,unit,normalization_note,record_status)
                        VALUES(?,'ALL',?,?,?,?,?,1,?,?,?,?,?,'有效')""", (trial_id, kind, f"{prefix}-{order:03d}", order, category, text,
                        spec.get("structured_field"), spec.get("operator"), spec.get("structured_value"), spec.get("unit"), "合作方提供的可计算条件" if spec else "仅保留原文，进入人工复核"))
                    counts[count_key] += 1
            counts["试验"] += 1
        connection.execute("UPDATE import_batches SET row_counts_json=? WHERE id=?", (json.dumps(counts, ensure_ascii=False), batch_id)); replace_trial_search(connection)
    return counts
